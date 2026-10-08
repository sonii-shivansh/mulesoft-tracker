import json
import re
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse, parse_qs, unquote

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data" / "jobs.json"
XLSX = ROOT / "MuleSoft_Job_Tracker.xlsx"
IST = ZoneInfo("Asia/Kolkata")
NOW = datetime.now(IST).strftime("%Y-%m-%d %H:%M IST")

HEADERS = {
    "User-Agent": "Mozilla/5.0 (compatible; MuleSoftJobTracker/1.0; +https://github.com/sonii-shivansh/mulesoft-tracker)"
}

PROFILE_TERMS = {
    "mulesoft": 22, "mule 4": 18, "anypoint": 15, "dataweave": 13,
    "raml": 10, "apikit": 8, "api-led": 8, "cloudhub": 8,
    "api manager": 7, "runtime manager": 5, "salesforce": 5,
    "sftp": 4, "soap": 4, "rest": 4, "munit": 5, "oauth": 3,
    "jwt": 3, "sql": 3, "database": 3, "anypoint mq": 3
}
POSITIVE_TITLE = (
    "mulesoft", "mule", "integration engineer", "integration developer",
    "api developer", "api engineer", "integration specialist"
)
NEGATIVE_TITLE = (
    "architect", "architecture", "manager", "director", "principal",
    "lead", "technical lead", "staff", "6+ years", "7+ years",
    "8+ years", "9+ years", "10+ years"
)

SEARCHES = [
    ("MuleSoft India jobs", '"MuleSoft Developer" India 2 years jobs'),
    ("MuleSoft Integration India", '"MuleSoft" "Integration" India jobs'),
    ("MuleSoft 2+ years", '"MuleSoft" "2+ years" India jobs'),
    ("MuleSoft 3 years", '"MuleSoft Developer" "3 years" India jobs'),
    ("Indeed MuleSoft", 'site:in.indeed.com "Mulesoft" India jobs'),
    ("Glassdoor MuleSoft", 'site:glassdoor.co.in/Job "MuleSoft Developer" India'),
    ("ZipRecruiter MuleSoft", 'site:ziprecruiter.in/jobs Mulesoft India'),
    ("Accenture MuleSoft", 'site:accenture.com/in-en/careers/jobdetails MuleSoft India'),
    ("Capgemini MuleSoft", 'site:careers.capgemini.com/job MuleSoft India'),
    ("Deloitte MuleSoft", 'site:deloitte.com "MuleSoft" India careers'),
    ("Cognizant MuleSoft", 'site:careers.cognizant.com/india-en/jobs MuleSoft'),
    ("PwC MuleSoft", 'site:linkedin.com/jobs "MuleSoft Developer" "PwC" India'),
    ("NTT DATA MuleSoft", 'site:linkedin.com/jobs "Mulesoft Developer" "NTT DATA" India'),
    ("Job boards MuleSoft", 'site:linkedin.com/jobs/view "MuleSoft" India "2+ years"'),
]

def canonical(url: str) -> str:
    if not url:
        return ""
    url = unquote(url.strip())
    parsed0 = urlparse(url)
    if parsed0.netloc.endswith("bing.com") and parsed0.path.startswith("/ck/a"):
        encoded = parse_qs(parsed0.query).get("u", [""])[0]
        if encoded.startswith("a1"):
            import base64
            try:
                padded = encoded[2:] + "=" * (-len(encoded[2:]) % 4)
                url = base64.b64decode(padded).decode("utf-8", errors="ignore")
            except Exception:
                return ""
    url = unquote(url.strip())
    parsed = urlparse(url)
    if parsed.netloc.endswith("google.com") and parsed.path == "/url":
        url = parse_qs(parsed.query).get("q", [""])[0]
        parsed = urlparse(url)
    if parsed.netloc.startswith("www."):
        parsed = parsed._replace(netloc=parsed.netloc[4:])
    parsed = parsed._replace(fragment="", query="")
    return parsed.geturl().rstrip("/")

def extract_bing(query):
    try:
        r = requests.get(
            "https://www.bing.com/search",
            params={"q": query, "count": 10},
            headers=HEADERS,
            timeout=20,
        )
        r.raise_for_status()
        soup = BeautifulSoup(r.text, "html.parser")
        out = []
        for item in soup.select("li.b_algo")[:10]:
            a = item.select_one("h2 a")
            if not a:
                continue
            href = canonical(a.get("href", ""))
            title = a.get_text(" ", strip=True)
            p = item.select_one(".b_caption p")
            snippet = p.get_text(" ", strip=True) if p else ""
            if href and title and urlparse(href).netloc.lower() not in {"bing.com", "www.bing.com"}:
                out.append({"url": href, "title": title, "snippet": snippet, "engine": "Bing"})
        return out
    except Exception as exc:
        print(f"SOURCE_ERROR=Bing | {type(exc).__name__}: {exc}")
        return []

def extract_ddg(query):
    try:
        r = requests.get(
            "https://html.duckduckgo.com/html/",
            params={"q": query},
            headers=HEADERS,
            timeout=20,
        )
        r.raise_for_status()
        soup = BeautifulSoup(r.text, "html.parser")
        out = []
        for a in soup.select("a.result__a")[:10]:
            href = canonical(a.get("href", ""))
            title = a.get_text(" ", strip=True)
            parent = a.find_parent("div", class_="result")
            snippet = ""
            if parent:
                s = parent.select_one(".result__snippet")
                snippet = s.get_text(" ", strip=True) if s else ""
            if href and title:
                out.append({"url": href, "title": title, "snippet": snippet})
        return out
    except Exception as exc:
        print(f"SOURCE_ERROR=DuckDuckGo | {type(exc).__name__}: {exc}")
        return []

def infer_company(title, url, text):
    host = urlparse(url).netloc.lower()
    if "accenture.com" in host:
        return "Accenture"
    if "capgemini.com" in host:
        return "Capgemini"
    if "deloitte.com" in host:
        return "Deloitte"
    if "cognizant.com" in host:
        return "Cognizant"
    if "pwc" in host or ("linkedin.com" in host and "pwc" in text.lower()):
        return "PwC Acceleration Center India"
    if "nttdata" in host or "ntt data" in text.lower():
        return "NTT DATA"
    if "linkedin.com" in host:
        parts = re.split(r"\s+at\s+|\s+-\s+", title, maxsplit=1)
        if len(parts) == 2:
            return parts[1].strip()
    return title.split(" - ")[-1].strip()[:80] or "Unknown"

def score_job(title, text):
    blob = (title + " " + text).lower()
    score = 0
    for term, points in PROFILE_TERMS.items():
        if term in blob:
            score += points
    if any(x in title.lower() for x in POSITIVE_TITLE):
        score += 10
    if any(x in title.lower() for x in NEGATIVE_TITLE):
        score -= 35
    years = re.findall(r"(\d+)\s*\+?\s*years?", blob)
    if years:
        minimums = [int(x) for x in years]
        if min(minimums) >= 6:
            score -= 45
        elif 2 <= min(minimums) <= 4:
            score += 15
    return max(0, min(100, score))

def relevant(title, text, url):
    blob = (title + " " + text + " " + url).lower()
    if not any(k in blob for k in (
        "mulesoft", "mule 4", "anypoint", "integration",
        "api developer", "api engineer"
    )):
        return False
    if any(k in title.lower() for k in (
        "architect", "manager", "director", "principal", "staff", "technical lead"
    )):
        return False
    years = [int(x) for x in re.findall(r"(\d+)\s*\+?\s*years?", blob)]
    if years and min(years) >= 6:
        return False
    return True

def make_note(title, text, score):
    blob = (title + " " + text).lower()
    hits = [term for term in PROFILE_TERMS if term in blob][:7]
    if not hits:
        return "Relevant MuleSoft/integration/API role; manual JD review recommended."
    return f"Matched profile skills: {', '.join(hits)}. Score {score}/100."

def load_jobs():
    if not DATA.exists():
        return []
    try:
        return json.loads(DATA.read_text(encoding="utf-8"))
    except Exception:
        return []

def save_jobs(jobs):
    DATA.parent.mkdir(parents=True, exist_ok=True)
    DATA.write_text(
        json.dumps(jobs, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )

def write_xlsx(jobs):
    # Rebuild from jobs.json every run so removed/stale rows cannot persist.
    wb = Workbook()
    ws = wb.active
    ws.title = "MuleSoft Jobs"
    ws.append([
        "Date & Time Found", "Company Name", "Job Title",
        "Location / Remote Status", "Match Score / Relevance Note",
        "Direct Application Link"
    ])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
    for job in jobs:
        ws.append([
            job["date_time_found"], job["company"], job["title"],
            job["location"],
            f'{job["score"]}/100 — {job["note"]}',
            job["url"],
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [22, 28, 42, 34, 75, 80]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(XLSX)

def main():
    jobs = load_jobs()
    known = {canonical(j.get("url", "")) for j in jobs}
    candidates = []

    source_results = 0
    for _, query in SEARCHES:
        items = extract_bing(query) + extract_ddg(query)
        source_results += len(items)
        for item in items:
            url = canonical(item["url"])
            if not url or url in known:
                continue
            text = item["title"] + " " + item.get("snippet", "")
            if not relevant(item["title"], text, url):
                continue
            score = score_job(item["title"], text)
            if score < 45:
                continue
            candidates.append({
                "date_time_found": NOW,
                "company": infer_company(item["title"], url, text),
                "title": item["title"],
                "location": "India / verify posting",
                "score": score,
                "note": make_note(item["title"], text, score),
                "url": url,
                "source": f"Public search result ({item.get('engine', 'unknown')})",
            })
        time.sleep(0.5)

    best = {}
    for job in candidates:
        key = canonical(job["url"])
        if key not in best or job["score"] > best[key]["score"]:
            best[key] = job

    new_jobs = sorted(
        best.values(),
        key=lambda x: (-x["score"], x["company"], x["title"])
    )
    jobs.extend(new_jobs)
    jobs.sort(key=lambda x: x.get("date_time_found", ""), reverse=True)
    save_jobs(jobs)
    write_xlsx(jobs)

    print(f"SEARCH_RESULTS={source_results}")
    print(f"NEW_JOBS={len(new_jobs)}")
    for j in new_jobs:
        print(f'- {j["company"]} | {j["title"]} | {j["score"]}/100 | {j["url"]}')

if __name__ == "__main__":
    main()
